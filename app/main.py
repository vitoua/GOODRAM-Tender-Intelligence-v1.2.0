from datetime import datetime,timedelta
from io import BytesIO
import asyncio,json,uuid
from pathlib import Path
from fastapi import FastAPI,Form,Request,HTTPException
from fastapi.responses import JSONResponse,StreamingResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from app.config import settings
from app.core import store,active,analytics
from app.sources import COUNTRIES
app=FastAPI(version='1.2.0');app.mount('/static',StaticFiles(directory='app/static'),name='static');tpl=Jinja2Templates(directory='app/templates');JOB={'running':False,'id':None,'progress':0,'message':'','error':''};LOCK=asyncio.Lock()
def default_date():return (datetime.now()-timedelta(days=30)).date().isoformat()
def rows(q='',country='',source='',active_only=1):return [x for x in store.data.values() if (not active_only or active(x)) and (not country or x.country==country) and (not source or x.source==source) and (not q or q.casefold() in (x.title+x.buyer+x.description+' '.join(x.cpv)).casefold())]
@app.get('/health')
def health():return {'status':'ok','version':'1.2.0','job':JOB,'tenders':len(store.data)}
@app.get('/')
def home(request:Request,q:str='',country:str='',source:str='',active_only:int=1):return tpl.TemplateResponse('index.html',{'request':request,'rows':sorted(rows(q,country,source,active_only),key=lambda x:x.deadline or datetime.max),'runs':store.runs,'countries':COUNTRIES,'default_date':default_date(),'q':q,'country':country,'source':source,'active_only':active_only,'last_update':store.last_update})
async def execute(job_id,selected,sources,start,min_value):
 async with LOCK:
  JOB.update(running=True,id=job_id,progress=2,message='Rozpoczynanie aktualizacji...',error='');results=[];tasks=[]
  plan=[]
  if 'ted' in sources and any(c!='UKR' for c in selected):plan.append(('ted',[c for c in selected if c!='UKR']))
  if 'prozorro' in sources and 'UKR' in selected:plan.append(('prozorro',['UKR']))
  if 'germany' in sources and 'DEU' in selected:plan.append(('germany',['DEU']))
  if 'orlen' in sources and 'POL' in selected:plan.append(('orlen',['POL']))
  if 'junta' in sources and 'ESP' in selected:plan.append(('junta',['ESP']))
  if 'base' in sources and 'PRT' in selected:plan.append(('base',['PRT']))
  for i,(name,scope) in enumerate(plan,1):
   JOB.update(progress=int((i-1)/max(1,len(plan))*90)+5,message=f'Pobieranie: {name.upper()}')
   try:
    module=__import__('app.connectors.'+('basept' if name=='base' else name),fromlist=['fetch']);data=await asyncio.wait_for(module.fetch(scope,start,min_value),120);results.append((name,data,None))
   except Exception as e:results.append((name,[],e))
  store.replace(results);JOB.update(running=False,progress=100,message='Aktualizacja zakończona')
@app.post('/refresh/start')
async def refresh_start(countries:list[str]=Form(default=[]),sources:list[str]=Form(default=[]),min_value_eur:float=Form(0),published_from:str=Form(default='')):
 if JOB['running']:return JSONResponse({'ok':False,'message':'Aktualizacja jest już uruchomiona.'},status_code=409)
 if not countries:return JSONResponse({'ok':False,'message':'Wybierz co najmniej jeden kraj.'},status_code=400)
 if not sources:return JSONResponse({'ok':False,'message':'Wybierz co najmniej jedno źródło.'},status_code=400)
 job=str(uuid.uuid4());asyncio.create_task(execute(job,countries,sources,datetime.fromisoformat(published_from or default_date()),min_value_eur));return {'ok':True,'job':job}
@app.get('/refresh/status')
def status():return {**JOB,'last_update':store.last_update.isoformat() if store.last_update else None,'runs':store.runs,'count':len(store.data)}
@app.get('/analytics')
def analytics_page(request:Request):return tpl.TemplateResponse('analytics.html',{'request':request,'data':analytics(list(store.data.values())),'count':len(store.data)})
@app.get('/cpv')
def cpv_page(request:Request):return tpl.TemplateResponse('cpv.html',{'request':request,'ted':json.loads(Path('data/ted_cpv.json').read_text()),'prozorro':json.loads(Path('data/prozorro_cpv.json').read_text())})
@app.post('/cpv/save')
def cpv_save(source:str=Form(),code:list[str]=Form(default=[]),name:list[str]=Form(default=[]),active_codes:list[str]=Form(default=[])):
 file=Path('data/prozorro_cpv.json' if source=='prozorro' else 'data/ted_cpv.json');key='name' if source=='prozorro' else 'name_pl';enabled=set(active_codes);file.write_text(json.dumps([{'code':c,key:n,'active':c in enabled} for c,n in zip(code,name)],ensure_ascii=False,indent=2));return JSONResponse({'ok':True})
@app.get('/keywords')
def kw(request:Request):return tpl.TemplateResponse('keywords.html',{'request':request,'groups':json.loads(Path('data/keywords.json').read_text())})
@app.get('/tenders/{tid:path}')
def detail(tid,request:Request):
 x=store.data.get(tid)
 if not x:raise HTTPException(404)
 return tpl.TemplateResponse('detail.html',{'request':request,'x':x})
@app.get('/export.xlsx')
def export(q:str='',country:str='',source:str=''):
 from openpyxl import Workbook
 from openpyxl.styles import Font,PatternFill
 wb=Workbook();ws=wb.active;ws.title='Przetargi';h=['Źródło','Kraj','Numer','Nazwa','Zamawiający','Publikacja','Termin','Wartość','Waluta','EUR','CPV','Link'];ws.append(h)
 for x in rows(q,country,source,0):ws.append([x.source,x.country,x.procurement_id,x.title,x.buyer,x.published,x.deadline,x.value,x.currency,x.value_eur,', '.join(x.cpv),x.url])
 for c in ws[1]:c.fill=PatternFill('solid',fgColor='E30613');c.font=Font(color='FFFFFF',bold=True)
 ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions;out=BytesIO();wb.save(out);out.seek(0);return StreamingResponse(out,media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':'attachment; filename="GOODRAM-Tender-Intelligence-v1.2.xlsx"'})
